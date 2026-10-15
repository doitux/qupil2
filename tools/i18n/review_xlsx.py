# Builds the review workbook (Excel) for translators and for checking the German changes.
# Usage: python3 -I review_xlsx.py <raw.json> <dir with entries.json from build.py>
import json, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from catalog import DROPPED  # noqa: E402

RAW, OUTDIR = sys.argv[1], sys.argv[2]
entries = json.load(open(os.path.join(OUTDIR, 'entries.json'), encoding='utf-8'))
raw = json.load(open(RAW, encoding='utf-8'))
raw_de = {}
for r in raw:
    if r['type'] not in ('obsolete', 'vanished'):
        raw_de.setdefault(r['src'], r['de'])
out = os.path.join(OUTDIR, 'Qupil2-Uebersetzungen.xlsx')

AREAS = {
    'app': 'App', 'common': 'Allgemein', 'weekday': 'Wochentage', 'menu': 'Menü', 'nav': 'Navigation',
    'today': 'Heute', 'schedule': 'Stundenplan', 'lesson': 'Unterricht', 'student': 'Schüler', 'archive': 'Archiv',
    'note': 'Notizen', 'format': 'Formatierung', 'piece': 'Werke', 'activity': 'Aktivitäten',
    'instruments': 'Instrumente', 'library': 'Notenbibliothek', 'event': 'Veranstaltungen',
    'candidates': 'Vorspielkandidaten', 'reminders': 'Erinnerungen', 'reminder': 'Erinnerungen',
    'birthday': 'Geburtstage', 'metronome': 'Metronom', 'audio': 'Audio', 'doc': 'Dokumente', 'import': 'Import',
    'backup': 'Backup', 'db': 'Datenbank', 'migration': 'Migration', 'sync': 'Synchronisation',
    'settings': 'Einstellungen', 'appearance': 'Darstellung', 'language': 'Sprache', 'more': 'Mehr',
    'notify': 'Benachrichtigungen', 'defaults': 'Voreinstellungen', 'about': 'Über Qupil',
}

F = 'Arial'
font = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
head_font = Font(name=F, size=10, bold=True, color='FFFFFF')
head_fill = PatternFill('solid', fgColor='6E3C5A')
fix_fill = PatternFill('solid', fgColor='FFF0B3')
corr_fill = PatternFill('solid', fgColor='FFD9D4')
new_fill = PatternFill('solid', fgColor='E3EDFF')
wrap = Alignment(wrap_text=True, vertical='top')
thin = Side(style='thin', color='E4E1E8')
grid = Border(bottom=thin)


def show(v):
    if v is None:
        return ''
    if isinstance(v, list):
        return 'Einzahl: %s\nMehrzahl: %s' % (v[0], v[1])
    return v


def kind(changes):
    text = ' '.join(changes)
    if 'KORREKTUR' in text:
        return 'Korrektur'
    if 'Rechtschreibung' in text or 'Tippfehler' in text:
        return 'Rechtschreibung'
    return 'Technisch'


def sheet(ws, headers, widths, rows, fills=None):
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 28
    for n, r in enumerate(rows):
        ws.append(r)
        for c in ws[ws.max_row]:
            c.font, c.alignment, c.border = font, wrap, grid
        if fills and fills[n]:
            col, fill = fills[n]
            ws.cell(row=ws.max_row, column=col).fill = fill
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()
ws = wb.active
ws.title = 'Strings'
rows, fills = [], []
for e in entries:
    changes = [c for c in e['changes'] if c != 'neu in Qupil 2']
    area = AREAS.get(e['id'].split('.')[0], e['id'].split('.')[0])
    rows.append([e['id'], area, show(e['en']), show(e['de']), show(e['fr']), show(e['es']),
                 show(e['de141']), show(e['old']), '; '.join(changes), 'ja' if e['new'] else '', e['comment'] or ''])
    k = kind(changes) if changes else None
    if e['new']:
        fills.append((10, new_fill))
    elif k == 'Korrektur':
        fills.append((9, corr_fill))
    elif k == 'Rechtschreibung':
        fills.append((9, fix_fill))
    else:
        fills.append(None)
sheet(ws, ['ID', 'Bereich', 'English', 'Deutsch', 'Français', 'Español', 'Deutsch 1.4.1', 'Englisch 1.4.1',
           'Änderung DE', 'Neu', 'Hinweis'], [30, 16, 38, 38, 38, 38, 34, 30, 34, 6, 30], rows, fills)

ws2 = wb.create_sheet('Deutsch geändert')
order = {'Korrektur': 0, 'Rechtschreibung': 1, 'Technisch': 2}
rows2 = []
for e in entries:
    changes = [c for c in e['changes'] if c != 'neu in Qupil 2']
    if not changes or e['new']:
        continue
    rows2.append([e['id'], kind(changes), show(e['de141']), show(e['de']), '; '.join(changes)])
rows2.sort(key=lambda r: (order[r[1]], r[0]))
fills2 = [(2, corr_fill) if r[1] == 'Korrektur' else (2, fix_fill) if r[1] == 'Rechtschreibung' else None for r in rows2]
sheet(ws2, ['ID', 'Art', 'Deutsch 1.4.1', 'Deutsch Qupil 2', 'Was geändert wurde'], [30, 16, 44, 44, 50], rows2, fills2)

ws3 = wb.create_sheet('Entfernt')
rows3 = [[src, raw_de.get(src, ''), why] for src, why in DROPPED.items()]
sheet(ws3, ['Text in 1.4.1 (Quelle)', 'Deutsch 1.4.1', 'Warum nicht übernommen'], [44, 44, 56], rows3)

ws4 = wb.create_sheet('Legende')
ws4.column_dimensions['A'].width = 26
ws4.column_dimensions['B'].width = 90
lines = [
    ('Qupil 2 – Übersetzungen zur Durchsicht', None),
    (None, None),
    ('Zahlen', None),
    ('Einträge gesamt', '=COUNTA(Strings!A:A)-1'),
    ('davon neu in Qupil 2', '=COUNTIF(Strings!J:J,"ja")'),
    ('Deutsch geändert', "=COUNTA('Deutsch geändert'!A:A)-1"),
    ('  davon Korrektur', "=COUNTIF('Deutsch geändert'!B:B,\"Korrektur\")"),
    ('  davon Rechtschreibung', "=COUNTIF('Deutsch geändert'!B:B,\"Rechtschreibung\")"),
    ('  davon technisch', "=COUNTIF('Deutsch geändert'!B:B,\"Technisch\")"),
    ('Nicht übernommen', '=COUNTA(Entfernt!A:A)-1'),
    (None, None),
    ('Blatt „Strings“', None),
    ('ID', 'Schlüssel im Code: qsTrId("…") in QML, qtTrId("…") in C++. Präfix = Bereich der App.'),
    ('English', 'Quellsprache und Standard. Fehlt ein Text in einer Übersetzung, zeigt Qupil den englischen.'),
    ('Deutsch', 'Wortlaut aus 1.4.1. Nur technisch bereinigt (Kürzel-&, „ - Qupil“ im Titel, HTML, harte Zeilenumbrüche), dazu Rechtschreibfehler und eine inhaltliche Korrektur. Alles einzeln im Blatt „Deutsch geändert“.'),
    ('Français / Español', 'Sie-Form (vous / usted). Französisch mit geschütztem Leerzeichen vor : ? ! ; und in « » – in Excel nicht sichtbar, in der App korrekt.'),
    ('Deutsch 1.4.1 / Englisch 1.4.1', 'Die Originaltexte aus qupil_de.ts von 1.4.1 (Englisch war dort die Quelle im Code).'),
    ('Änderung DE', 'Gelb = Rechtschreibung, Rot = inhaltliche Korrektur, leer = unverändert oder nur technisch.'),
    ('Neu', 'Blau „ja“ = Text gibt es erst in Qupil 2 (Sync, Backup, Darstellung, Sprache, mobile Bedienung).'),
    ('Einzahl / Mehrzahl', 'Pluraltexte stehen zweizeilig. %n wird durch die Zahl ersetzt.'),
    ('%1, %2 …', 'Platzhalter, werden zur Laufzeit ersetzt. Reihenfolge darf sich je Sprache ändern.'),
    (None, None),
    ('Rückgängig machen', 'Soll eine Änderung am Deutschen zurück: in tools/catalog.py beim Eintrag de="…" mit dem alten Text setzen und tools/build.py laufen lassen.'),
]
for a, b in lines:
    ws4.append([a, b])
    r = ws4.max_row
    ws4.cell(row=r, column=1).font = bold if a and (b is None or r in (1,)) else font
    ws4.cell(row=r, column=2).font = font
    ws4.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical='top')
    ws4.cell(row=r, column=1).alignment = Alignment(vertical='top')
ws4['A1'].font = Font(name=F, size=14, bold=True, color='6E3C5A')
for r in range(4, 11):
    ws4.cell(row=r, column=2).alignment = Alignment(horizontal='left')
wb.move_sheet('Legende', offset=-3)
wb.active = 0
wb.save(out)
print('ok', len(rows), len(rows2), len(rows3))
